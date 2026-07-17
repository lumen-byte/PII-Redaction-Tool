
import sys
import json
import docx
from redactor import PIIRedactor


def redact_paragraph(paragraph, redactor: PIIRedactor):
    full_text = paragraph.text
    if not full_text or not full_text.strip():
        return
    redacted = redactor.redact_text(full_text)
    if redacted == full_text:
        return
    if not paragraph.runs:
        return
    paragraph.runs[0].text = redacted
    for run in paragraph.runs[1:]:
        run.text = ""


def redact_document(input_path: str, output_path: str, redactor: PIIRedactor):
    document = docx.Document(input_path)

    for paragraph in document.paragraphs:
        redact_paragraph(paragraph, redactor)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    redact_paragraph(paragraph, redactor)

    document.save(output_path)


def main():
    if len(sys.argv) < 3:
        print("Usage: python3 redact_docx.py input.docx output_redacted.docx [audit_log.json]")
        sys.exit(1)

    input_path = sys.argv[1]
    output_path = sys.argv[2]
    audit_path = sys.argv[3] if len(sys.argv) > 3 else "audit_log.json"

    redactor = PIIRedactor()
    redact_document(input_path, output_path, redactor)

    audit_log = redactor.get_audit_log()
    with open(audit_path, "w") as f:
        json.dump(audit_log, f, indent=2)

    mapping_path = audit_path.replace(".json", "_mapping.json")
    with open(mapping_path, "w") as f:
        json.dump(redactor.get_mapping(), f, indent=2)

    counts = {}
    for row in audit_log:
        counts[row["type"]] = counts.get(row["type"], 0) + 1

    print(f"Redacted document written to: {output_path}")
    print(f"Audit log ({len(audit_log)} replacements) written to: {audit_path}")
    print(f"Entity mapping written to: {mapping_path}")
    print("\nReplacements by category:")
    for label, n in sorted(counts.items(), key=lambda x: -x[1]):
        print(f"  {label:<14}{n}")


if __name__ == "__main__":
    main()
